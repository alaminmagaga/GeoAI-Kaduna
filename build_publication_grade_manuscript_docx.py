from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


WORKSPACE = Path(__file__).resolve().parent
OUT = WORKSPACE / "kaduna_publication_grade_outputs"
FIG = OUT / "figures"
TABLES = OUT / "tables"
DOCX = WORKSPACE / "GeoAI_Kaduna_Publication_Grade_Manuscript_With_Clear_Figures.docx"


TITLE = "GeoAI-Based Detection of Urban Growth Hotspots for Climate-Resilient Urban Planning in Kaduna Metropolis, Nigeria"


def set_portrait(section):
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)


def set_landscape(section):
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.42)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table_from_df(doc, df, caption, widths=None, font_size=8):
    p = doc.add_paragraph(caption)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.bold = True
        r.font.size = Pt(10)

    table = doc.add_table(rows=1, cols=len(df.columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, col in enumerate(df.columns):
        set_cell_text(hdr[i], col, bold=True, size=font_size)
        shade_cell(hdr[i], "D9EAF7")

    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if pd.isna(value):
                value = "-"
            set_cell_text(cells[i], value, size=font_size)

    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)


def add_figure_page(doc, image_name, caption, width=9.7):
    path = FIG / image_name
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_landscape(section)
    p = doc.add_paragraph(caption)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.bold = True
        r.font.size = Pt(11)
    img_p = doc.add_paragraph()
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = img_p.add_run()
    run.add_picture(str(path), width=Inches(width))
    note = doc.add_paragraph(f"Source figure: kaduna_publication_grade_outputs/figures/{image_name}")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in note.runs:
        r.italic = True
        r.font.size = Pt(7)
    portrait = doc.add_section(WD_SECTION.NEW_PAGE)
    set_portrait(portrait)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(31, 78, 121)
    return p


def add_para(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08
    return p


def build():
    stats = pd.read_csv(TABLES / "table_1_builtup_statistics.csv")
    stats_doc = stats.copy()
    stats_doc.columns = ["Year", "Built-up surface (sq. km)", "Change (sq. km)", "Percentage change"]
    stats_doc["Built-up surface (sq. km)"] = stats_doc["Built-up surface (sq. km)"].map(lambda x: f"{x:.1f}")
    stats_doc["Change (sq. km)"] = stats_doc["Change (sq. km)"].map(lambda x: "-" if pd.isna(x) else f"{x:.1f}")
    stats_doc["Percentage change"] = stats_doc["Percentage change"].map(lambda x: "-" if pd.isna(x) else f"{x:.1f}%")

    hotspots = pd.read_csv(TABLES / "table_2_top_hotspots_with_driver_metrics.csv")
    hot_doc = hotspots.rename(
        columns={
            "rank": "Rank",
            "lga": "LGA",
            "latitude": "Lat.",
            "longitude": "Lon.",
            "growth_sqkm_2000_2025": "Growth 2000-2025 (sq. km)",
            "distance_to_major_road_km": "Road dist. (km)",
            "distance_to_2000_builtup_core_km": "Core dist. (km)",
            "distance_to_river_km": "River dist. (km)",
            "planning_interpretation": "Planning interpretation",
        }
    )
    hot_doc = hot_doc[
        [
            "Rank",
            "LGA",
            "Lat.",
            "Lon.",
            "Growth 2000-2025 (sq. km)",
            "Road dist. (km)",
            "Core dist. (km)",
            "River dist. (km)",
            "Planning interpretation",
        ]
    ].copy()
    for col in ["Lat.", "Lon."]:
        hot_doc[col] = hot_doc[col].map(lambda x: f"{x:.5f}")
    hot_doc["Growth 2000-2025 (sq. km)"] = hot_doc["Growth 2000-2025 (sq. km)"].map(lambda x: f"{x:.3f}")

    doc = Document()
    set_portrait(doc.sections[0])
    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(TITLE)
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(31, 78, 121)
    meta = doc.add_paragraph("Alamin Musa Magaga | YandyTech Community / Nest Africa AI Innovation Lab | Nigeria | alaminmusamagaga@gmail.com")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading(doc, "Abstract")
    add_para(
        doc,
        "Rapid urban expansion is reshaping many African cities and increasing pressure on drainage, infrastructure, green spaces, and public services. This study applies a GeoAI-enabled Earth Observation workflow to detect urban growth hotspots in Kaduna Metropolis, Nigeria, from 2000 to 2025. The metropolitan area was defined using four Local Government Areas: Kaduna North, Kaduna South, Chikun, and Igabi. Built-up surface estimates were derived from the Global Human Settlement Layer GHS-BUILT-S P2023A product, while Landsat imagery, OpenStreetMap roads, water features, and administrative boundaries were used for visual and spatial context. Results show that built-up surface increased from 62.1 sq. km in 2000 to 129.7 sq. km in 2025, meaning the built-up footprint more than doubled over the 25-year period. Grid-based analysis identifies the strongest growth cells in Chikun, Igabi, Kaduna South, and Kaduna North, with coordinates and distance metrics linking hotspots to roads, the 2000 built-up core, and river or water features. The study demonstrates a reproducible open-data workflow for climate-resilient urban planning by showing where rapid expansion may require priority attention for drainage, infrastructure, service delivery, green-space protection, and field validation.",
    )
    add_para(doc, "Keywords: GeoAI; urban growth; GHSL; Kaduna Metropolis; climate resilience; Earth Observation; hotspot analysis")

    add_heading(doc, "1. Introduction")
    add_para(doc, "Urban growth is a major form of environmental and socioeconomic transformation in African cities. As built-up surfaces replace vegetation, agricultural land, and open spaces, cities may gain housing and economic activity while also increasing runoff, heat exposure, drainage stress, and infrastructure demand.")
    add_para(doc, "This paper frames Kaduna's urban growth as a climate-resilience planning problem. The goal is not only to measure how much the city expanded, but to identify where growth is concentrated so that planning agencies can prioritize drainage, transport, services, green-space protection, and field validation.")

    add_heading(doc, "2. Study Area")
    add_para(doc, "Kaduna Metropolis was extracted from Kaduna State using four FAO GAUL level-2 administrative units: Kaduna North, Kaduna South, Chikun, and Igabi. Kaduna North and Kaduna South represent the established urban core, while Chikun and Igabi capture major peri-urban expansion zones.")
    add_figure_page(doc, "figure_1_study_area_inset.png", "Figure 1. Kaduna Metropolitan Area within Kaduna State.")
    add_figure_page(doc, "figure_2_metropolitan_context_roads_river_core.png", "Figure 2. Metropolitan context showing LGAs, major roads, water features, and the 2000 built-up core.")

    add_heading(doc, "3. Data and Methods")
    add_para(doc, "The primary evidence layer was GHSL GHS-BUILT-S P2023A for 2000, 2005, 2010, 2015, 2020, and 2025. GHSL was selected because it provides a consistent multi-temporal built-up surface product and avoids the temporal bias observed when a classifier trained on a modern reference year is applied backward to older Landsat imagery.")
    add_para(doc, "Built-up growth was calculated as the 2025 built-up surface minus the 2000 built-up surface, with negative values set to zero. The study area was divided into regular grid cells, and each cell was ranked by total built-up surface increase. The ranked cells are treated as planning-oriented growth-intensity hotspots, not as formal inferential statistical clusters.")
    add_para(doc, "Spatial driver context was assessed by overlaying growth cells with OSM major roads, river and water features, LGA boundaries, and the 2000 built-up core. For each top hotspot cell, distances to the nearest major road, nearest river or water feature, and baseline built-up core were calculated.")

    add_heading(doc, "4. Results")
    add_heading(doc, "4.1 Built-Up Surface Trend", level=2)
    add_para(doc, "Built-up surface increased from 62.1 sq. km in 2000 to 129.7 sq. km in 2025. The largest interval gains occurred from 2005 to 2010 and from 2010 to 2015, while growth continued at a lower rate after 2015.")
    add_table_from_df(doc, stats_doc, "Table 1. GHSL built-up surface statistics for Kaduna Metropolis.", widths=[0.9, 1.7, 1.4, 1.5])
    add_figure_page(doc, "figure_3_builtup_surface_trend.png", "Figure 3. Built-up surface trend in Kaduna Metropolis, 2000-2025.")
    add_figure_page(doc, "figure_4_builtup_expansion_by_period.png", "Figure 4. Built-up expansion by period in Kaduna Metropolis.")
    add_figure_page(doc, "figure_5_growth_intensity_2000_2025.png", "Figure 5. Built-up growth intensity from 2000 to 2025.")

    add_heading(doc, "4.2 Growth Hotspots and Spatial Context", level=2)
    add_para(doc, "The strongest growth cells are concentrated around the central metropolitan belt and adjacent peri-urban areas. Chikun and Igabi contain several high-ranking cells, while Kaduna South and Kaduna North remain important because they contain central growth and edge-expansion locations.")
    add_figure_page(doc, "figure_6_urban_growth_hotspots.png", "Figure 6. Urban growth intensity hotspots, 2000-2025.")
    add_figure_page(doc, "figure_7_hotspots_spatial_drivers_overlay.png", "Figure 7. Hotspot and spatial driver context overlay.")
    add_table_from_df(doc, hot_doc, "Table 2. Top built-up growth hotspot cells with coordinates and driver metrics.", widths=[0.45, 0.85, 0.65, 0.65, 0.85, 0.65, 0.65, 0.65, 2.7], font_size=7)
    add_figure_page(doc, "figure_8_before_after_hotspot_panels.png", "Figure 8. Satellite context for top growth hotspots, 2010 and 2025.")

    add_heading(doc, "5. Discussion")
    add_para(doc, "The results indicate substantial and spatially uneven urban expansion in Kaduna Metropolis. Growth is strongest around the established metropolitan belt and peri-urban extension zones, especially in Chikun and Igabi. This pattern is consistent with outward metropolitan expansion beyond the older core LGAs.")
    add_para(doc, "The climate-resilience relevance comes from the physical consequences of built-up expansion. More impervious surface can increase runoff, reduce infiltration, raise local heat exposure, and place pressure on drainage and public services. Hotspot mapping therefore provides a practical way to identify where adaptation-oriented planning attention should be prioritized.")
    add_para(doc, "The earlier Random Forest LULC experiment produced implausible historical built-up results, including a false built-up decline. The corrected workflow uses GHSL as the primary long-term evidence layer and uses Landsat and contextual data for interpretation. This correction strengthens the publication value of the analysis because it explicitly handles visual QA and plausibility checking.")

    add_heading(doc, "6. Limitations")
    add_para(doc, "GHSL is a generalized built-up surface product and may miss small structures or fine local changes. The hotspot grid identifies high-growth cells and spatial associations, but it does not prove causal drivers. Coordinates are grid-cell centroids and should be used for field orientation rather than parcel-level decisions.")

    add_heading(doc, "7. Conclusion")
    add_para(doc, "This study provides a reproducible GeoAI workflow for detecting urban growth hotspots in Kaduna Metropolis using open Earth Observation and geospatial data. The built-up surface more than doubled between 2000 and 2025, with growth concentrated in Chikun, Igabi, Kaduna South, and Kaduna North. The outputs support climate-resilient planning by identifying priority areas for drainage assessment, infrastructure planning, green-space protection, service delivery, and field validation.")

    add_heading(doc, "References")
    add_para(doc, "European Commission, Joint Research Centre. Global Human Settlement Layer GHS-BUILT-S P2023A. Google Earth Engine dataset: JRC/GHSL/P2023A/GHS_BUILT_S.")
    add_para(doc, "FAO. GAUL 2015 administrative boundaries. Google Earth Engine dataset: FAO/GAUL/2015.")
    add_para(doc, "OpenStreetMap contributors. Road and water feature data accessed through OSMnx.")
    add_para(doc, "United States Geological Survey. Landsat Collection 2 Level 2 surface reflectance imagery.")

    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    build()
