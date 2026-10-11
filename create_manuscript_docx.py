from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


WORKSPACE = Path(__file__).resolve().parent
MD_PATH = WORKSPACE / "GeoAI_Kaduna_Climate_Resilient_Urban_Growth_Manuscript.md"
DOCX_PATH = WORKSPACE / "GeoAI_Kaduna_Climate_Resilient_Urban_Growth_Manuscript.docx"


FIGURES = {
    "kaduna_publication_outputs/figure_1_kaduna_state_with_metropolitan_lgas.png": "Figure 1. Kaduna Metropolitan Area within Kaduna State.",
    "kaduna_publication_outputs/figure_2_verified_kaduna_metropolitan_lgas.png": "Figure 2. Verified Kaduna metropolitan study area.",
    "kaduna_publication_outputs/figure_3_ghsl_builtup_growth_trend.png": "Figure 3. GHSL built-up surface growth trend.",
    "kaduna_publication_outputs/figure_4_ghsl_built_surface_panel.png": "Figure 4. Multi-year GHSL built-up surface map panel.",
    "kaduna_publication_outputs/figure_5_satellite_builtup_comparison_2000_2025.png": "Figure 5. Satellite imagery with dense built-up overlay.",
    "kaduna_publication_outputs/figure_6_builtup_growth_intensity_2000_2025.png": "Figure 6. Built-up growth intensity from 2000 to 2025.",
    "kaduna_publication_outputs/figure_7_growth_hotspot_map_labeled.png": "Figure 7. Urban growth hotspot zones from 2000 to 2025.",
    "kaduna_publication_outputs/figure_8_top_growth_locations_table.png": "Figure 8. Top growth locations table.",
    "kaduna_publication_outputs/figure_9_recent_satellite_landcover_context_2025.png": "Figure 9. Recent satellite and land-cover context for 2025.",
}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_markdown_table(document, lines):
    rows = []
    for line in lines:
        stripped = line.strip()
        if not stripped.startswith("|"):
            continue
        parts = [part.strip() for part in stripped.strip("|").split("|")]
        if all(set(part) <= {"-", ":"} for part in parts):
            continue
        rows.append(parts)

    if not rows:
        return

    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = value
            if r_idx == 0:
                set_cell_shading(cell, "D9EAF7")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True


def add_paragraph_with_basic_bold(document, text):
    paragraph = document.add_paragraph()
    # Basic support for **bold** spans.
    parts = text.split("**")
    for idx, part in enumerate(parts):
        run = paragraph.add_run(part)
        if idx % 2 == 1:
            run.bold = True
    return paragraph


def add_figure(document, rel_path, caption):
    path = WORKSPACE / rel_path
    if not path.exists():
        document.add_paragraph(f"[Missing figure: {rel_path}]")
        return

    document.add_picture(str(path), width=Inches(6.3))
    last = document.paragraphs[-1]
    last.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_para = document.add_paragraph(caption)
    caption_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in caption_para.runs:
        run.italic = True


def build_docx():
    text = MD_PATH.read_text(encoding="utf-8")
    lines = text.splitlines()
    document = Document()

    styles = document.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)

    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if not line:
            i += 1
            continue

        if line.startswith("# "):
            title = line[2:].strip()
            paragraph = document.add_heading(title, level=0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        if line.startswith("## "):
            document.add_heading(line[3:].strip(), level=1)
            i += 1
            continue

        if line.startswith("### "):
            document.add_heading(line[4:].strip(), level=2)
            i += 1
            continue

        if line.startswith("!["):
            start = line.find("](")
            end = line.rfind(")")
            caption = line[2:start]
            rel_path = line[start + 2 : end]
            add_figure(document, rel_path, caption)
            i += 1
            continue

        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_markdown_table(document, table_lines)
            continue

        if line.startswith("- "):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
            i += 1
            continue

        add_paragraph_with_basic_bold(document, line)
        i += 1

    document.save(DOCX_PATH)
    print(f"Wrote {DOCX_PATH}")


if __name__ == "__main__":
    build_docx()
