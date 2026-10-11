from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


WORKSPACE = Path(__file__).resolve().parent
MD_PATH = WORKSPACE / "GeoAI_Kaduna_Climate_Resilient_Urban_Growth_Manuscript.md"
DOCX_PATH = WORKSPACE / "GeoAI_Kaduna_Climate_Resilient_Urban_Growth_Manuscript_Clear_Figures.docx"


def set_portrait(section):
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)


def set_landscape(section):
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_markdown_table(document, lines):
    rows = []
    for line in lines:
        stripped = line.strip()
        if not stripped.startswith("|"):
            continue
        parts = [part.strip() for part in stripped.strip("|").split("|")]
        if all(set(part) <= {"-", ":"} for part in parts):
            continue
        rows.append(parts)

    if not rows:
        return

    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = value
            if r_idx == 0:
                set_cell_shading(cell, "D9EAF7")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True


def add_paragraph_with_basic_bold(document, text):
    paragraph = document.add_paragraph()
    parts = text.split("**")
    for idx, part in enumerate(parts):
        run = paragraph.add_run(part)
        if idx % 2 == 1:
            run.bold = True
    return paragraph


def add_clear_figure_page(document, rel_path, caption):
    path = WORKSPACE / rel_path
    if not path.exists():
        document.add_paragraph(f"[Missing figure: {rel_path}]")
        return

    landscape = document.add_section(WD_SECTION.NEW_PAGE)
    set_landscape(landscape)

    caption_para = document.add_paragraph(caption)
    caption_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in caption_para.runs:
        run.bold = True

    picture_para = document.add_paragraph()
    picture_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = picture_para.add_run()
    run.add_picture(str(path), width=Inches(9.85))

    source_para = document.add_paragraph(f"Image file: {rel_path}")
    source_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in source_para.runs:
        run.italic = True
        run.font.size = Pt(8)

    portrait = document.add_section(WD_SECTION.NEW_PAGE)
    set_portrait(portrait)


def build_docx():
    text = MD_PATH.read_text(encoding="utf-8")
    lines = text.splitlines()
    document = Document()

    styles = document.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)

    set_portrait(document.sections[0])

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if not line:
            i += 1
            continue

        if line.startswith("# "):
            paragraph = document.add_heading(line[2:].strip(), level=0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        if line.startswith("## "):
            document.add_heading(line[3:].strip(), level=1)
            i += 1
            continue

        if line.startswith("### "):
            document.add_heading(line[4:].strip(), level=2)
            i += 1
            continue

        if line.startswith("!["):
            start = line.find("](")
            end = line.rfind(")")
            caption = line[2:start]
            rel_path = line[start + 2 : end]
            add_clear_figure_page(document, rel_path, caption)
            i += 1
            continue

        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_markdown_table(document, table_lines)
            continue

        if line.startswith("- "):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
            i += 1
            continue

        add_paragraph_with_basic_bold(document, line)
        i += 1

    document.save(DOCX_PATH)
    print(f"Wrote {DOCX_PATH}")


if __name__ == "__main__":
    build_docx()
